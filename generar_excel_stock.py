"""Excel de stock desde bc.[Item Ledger Entry] (solo Cajas pendientes = 1).

Cajas pendientes = Remaining Quantity. Solo filas Open=1 y Remaining Quantity=1.
Stock disponible = Kilos de esa misma linea ILE. No se usa otra tabla para el stock.
"""
from __future__ import annotations

import datetime as dt
import os
from pathlib import Path

import pymssql
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent
BIOMASA_ENV = Path(r"C:\Users\ACZ\.codex\CALCULO_BIOMASA\.env")
OUT_DIR = ROOT / "Reports"

SQL_STOCK = """
SELECT
    ile.[Location Code] AS almacen,
    ile.[Item No.] AS item_no,
    ile.[Description] AS producto,
    CAST(ile.[Kilos] AS float) AS stock_kg,
    CAST(ile.[Remaining Quantity] AS float) AS cajas_pendientes,
    CAST(ile.[Fecha despesque] AS date) AS fecha_despesque,
    CAST(ile.[Fecha empaque] AS date) AS fecha_empaque,
    CAST(ile.[Expiration Date] AS date) AS fecha_caducidad,
    ile.[Lot No.] AS lote
FROM bc.[Item Ledger Entry] AS ile
WHERE ile.[Open] = 1
  AND ile.[Remaining Quantity] = 1
"""

HEADERS = [
    "Almacen",
    "Item",
    "Producto",
    "Stock disponible (kg)",
    "Fecha despesque",
    "Fecha empaque",
    "Fecha caducidad",
    "Cajas pendientes",
    "Dias a caducidad",
    "Estado caducidad",
]

FILLS = {
    "Caducado": PatternFill("solid", fgColor="FFC7CE"),
    "Proximo a caducar": PatternFill("solid", fgColor="F4B183"),
    "En seguimiento": PatternFill("solid", fgColor="FFE699"),
    "Correcto": PatternFill("solid", fgColor="C6EFCE"),
    "Sin fecha": PatternFill("solid", fgColor="D9D9D9"),
}
FONTS = {
    "Caducado": Font(color="9C0006", bold=True),
    "Proximo a caducar": Font(color="C65911", bold=True),
    "En seguimiento": Font(color="806000", bold=True),
    "Correcto": Font(color="006100", bold=True),
    "Sin fecha": Font(color="595959"),
}


def load_env(path: Path) -> None:
    if not path.exists():
        return
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        if "=" in line:
            key, value = line.split("=", 1)
            key, value = key.strip(), value.strip().strip('"').strip("'")
            if key and key not in os.environ:
                os.environ[key] = value
            continue
        if ":" not in line:
            continue
        label, value = line.split(":", 1)
        mapped = {
            "server name": "BC_SERVER",
            "database": "BC_DATABASE",
            "user": "BC_USER",
            "password": "BC_PASSWORD",
        }.get(label.strip().lower())
        if mapped and mapped not in os.environ and value.strip():
            os.environ[mapped] = value.strip()


def limpia_fecha(v) -> dt.date | None:
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        v = v.date()
    if not isinstance(v, dt.date):
        return None
    if v.year < 1900:
        return None
    return v


def estado_caducidad(exp: dt.date | None, today: dt.date) -> tuple[str, int | None]:
    exp = limpia_fecha(exp)
    if exp is None:
        return "Sin fecha", None
    dias = (exp - today).days
    if dias < 0:
        return "Caducado", dias
    if dias <= 7:
        return "Proximo a caducar", dias
    if dias <= 15:
        return "En seguimiento", dias
    return "Correcto", dias


def fetch_rows() -> list[tuple]:
    load_env(BIOMASA_ENV)
    load_env(ROOT / ".env")
    missing = [k for k in ("BC_SERVER", "BC_DATABASE", "BC_USER", "BC_PASSWORD") if not os.environ.get(k)]
    if missing:
        raise SystemExit("Faltan credenciales BC SQL: " + ", ".join(missing))
    conn = pymssql.connect(
        server=os.environ["BC_SERVER"],
        user=os.environ["BC_USER"],
        password=os.environ["BC_PASSWORD"],
        database=os.environ["BC_DATABASE"],
        login_timeout=60,
        timeout=600,
    )
    cur = conn.cursor()
    cur.execute(SQL_STOCK)
    rows = cur.fetchall()
    conn.close()
    return rows


def style_header(ws) -> None:
    fill = PatternFill("solid", fgColor="003051")
    font = Font(color="FFFFFF", bold=True, name="Calibri")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 22


def write_stock(ws, data: list[list], today: dt.date) -> None:
    ws.append(HEADERS)
    for row in data:
        ws.append(row)
    last = max(2, 1 + len(data))
    style_header(ws)
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    for r in range(2, last + 1):
        estado = ws.cell(r, 10).value
        fill = FILLS.get(str(estado or "Sin fecha"))
        font = FONTS.get(str(estado or "Sin fecha"))
        for c in range(1, 11):
            cell = ws.cell(r, c)
            cell.border = thin
            cell.font = Font(name="Calibri", size=10, color=(font.color.rgb if font and c == 10 else "000000"), bold=(c == 10))
            if c == 10 and fill:
                cell.fill = fill
                if font:
                    cell.font = font
        ws.cell(r, 4).number_format = "#,##0.00"
        ws.cell(r, 5).number_format = "DD/MM/YYYY"
        ws.cell(r, 6).number_format = "DD/MM/YYYY"
        ws.cell(r, 7).number_format = "DD/MM/YYYY"
        ws.cell(r, 8).number_format = "0"
        ws.cell(r, 9).number_format = "0"
    widths = [14, 14, 44, 20, 16, 16, 16, 16, 16, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    tab = Table(displayName="StockILE", ref=f"A1:J{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False


def write_resumen(ws, data: list[list], today: dt.date, n: int) -> None:
    ws["A1"] = "Stock disponible — Item Ledger Entry"
    ws["A1"].font = Font(name="Calibri", size=18, bold=True, color="003051")
    ws["A2"] = (
        "Fuente unica: bc.[Item Ledger Entry]. "
        "Filtro: Open = 1 y Remaining Quantity (Cajas pendientes) = 1. "
        "No se usa Inventory ni otra tabla para el stock."
    )
    ws["A2"].font = Font(name="Calibri", size=10, color="595959")
    ws.merge_cells("A2:F2")
    ws["A3"] = f"Generado {today.strftime('%d/%m/%Y')}  |  {n} registros"
    ws["A3"].font = Font(name="Calibri", size=11, bold=True)

    ws["A5"] = "Estado caducidad"
    ws["B5"] = "Registros"
    ws["C5"] = "Kg"
    for cell in (ws["A5"], ws["B5"], ws["C5"]):
        cell.fill = PatternFill("solid", fgColor="003051")
        cell.font = Font(color="FFFFFF", bold=True)

    counts: dict[str, list[float]] = {}
    for row in data:
        est = str(row[9])
        kg = float(row[3] or 0)
        if est not in counts:
            counts[est] = [0, 0.0]
        counts[est][0] += 1
        counts[est][1] += kg

    order = ["Caducado", "Proximo a caducar", "En seguimiento", "Correcto", "Sin fecha"]
    r = 6
    for est in order:
        if est not in counts:
            continue
        ws.cell(r, 1, est)
        ws.cell(r, 2, counts[est][0])
        ws.cell(r, 3, round(counts[est][1], 2))
        ws.cell(r, 3).number_format = "#,##0.00"
        if est in FILLS:
            ws.cell(r, 1).fill = FILLS[est]
            ws.cell(r, 1).font = FONTS[est]
        r += 1

    ws["A12"] = "Leyenda"
    ws["A12"].font = Font(bold=True, size=12, color="003051")
    legend = [
        ("Caducado", "Fecha de caducidad anterior a hoy"),
        ("Proximo a caducar", "7 dias o menos"),
        ("En seguimiento", "Entre 8 y 15 dias"),
        ("Correcto", "Mas de 15 dias"),
    ]
    rr = 13
    for name, txt in legend:
        ws.cell(rr, 1, name).fill = FILLS[name]
        ws.cell(rr, 1).font = FONTS[name]
        ws.cell(rr, 2, txt)
        rr += 1
    ws["A18"] = "En la hoja Stock usa las flechas de filtro de Almacen, Producto y Estado caducidad."
    ws["A18"].font = Font(italic=True, color="595959")
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 44
    ws.column_dimensions["C"].width = 16
    ws.sheet_view.showGridLines = False


def main() -> None:
    today = dt.date.today()
    print("Consultando Item Ledger Entry (Open=1, Remaining Quantity=1)...")
    raw = fetch_rows()
    print(f"  {len(raw)} filas")
    data: list[list] = []
    for almacen, item_no, producto, kg, cajas, fdesp, femp, fcad, _lote in raw:
        estado, dias = estado_caducidad(fcad, today)
        data.append(
            [
                almacen,
                item_no,
                producto,
                None if kg is None else round(float(kg), 2),
                limpia_fecha(fdesp),
                limpia_fecha(femp),
                limpia_fecha(fcad),
                int(cajas or 0),
                dias,
                estado,
            ]
        )

    wb = Workbook()
    ws_r = wb.active
    ws_r.title = "Resumen"
    write_resumen(ws_r, data, today, len(data))
    ws = wb.create_sheet("Stock", 1)
    write_stock(ws, data, today)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    dest = OUT_DIR / f"stock_ile_{today.strftime('%Y%m%d')}.xlsx"
    wb.save(dest)
    print("Guardado", dest)


if __name__ == "__main__":
    main()
